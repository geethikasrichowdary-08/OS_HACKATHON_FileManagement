#!/usr/bin/env python3
"""Generate the Team G-3 Mini Linux Shell hackathon presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent / "minish_team_g3.pptx"
INK = RGBColor(25, 43, 48)
PAPER = RGBColor(247, 246, 240)
TEAL = RGBColor(13, 113, 103)
MINT = RGBColor(208, 231, 220)
ORANGE = RGBColor(218, 105, 57)
MUTED = RGBColor(91, 108, 107)
WHITE = RGBColor(255, 255, 255)
PALE = RGBColor(232, 237, 230)


def textbox(slide, x, y, w, h, text, size=18, color=INK, bold=False,
            font="Aptos"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.space_after = Pt(8)
        paragraph.font.name = font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, color, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.08
    return shape


def new_slide(deck, number, title, kicker):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    rect(slide, 0, 0, 0.16, 7.5, TEAL)
    textbox(slide, 0.55, 0.35, 5, 0.25, kicker.upper(), 10, TEAL, True)
    textbox(slide, 0.55, 0.7, 12.1, 0.6, title, 28, INK, True, "Aptos Display")
    rect(slide, 0.55, 1.42, 12.2, 0.02, PALE)
    textbox(slide, 11.9, 7.05, 0.8, 0.25, f"G-3  /  {number:02d}", 9, MUTED, True)
    return slide


def split_slide(deck, number, title, kicker, items, panel_title, panel_lines):
    slide = new_slide(deck, number, title, kicker)
    y = 1.82
    for heading, detail in items:
        textbox(slide, 0.7, y, 5.6, 0.34, heading, 17, INK, True)
        textbox(slide, 0.7, y + 0.37, 5.55, 0.65, detail, 13, MUTED)
        y += 1.25
    rect(slide, 6.75, 1.8, 5.6, 4.85, MINT, True)
    textbox(slide, 7.1, 2.15, 4.9, 0.48, panel_title, 20, TEAL, True)
    textbox(slide, 7.1, 2.88, 4.8, 3.35, "\n".join(panel_lines), 15, INK)
    return slide


def build_deck():
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    deck.core_properties.title = "Mini Linux Shell with Background Job Management"
    deck.core_properties.subject = "OS Hackathon project, Team G-3"
    deck.core_properties.author = "Team G-3"

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = INK
    rect(slide, 0, 0, 0.22, 7.5, ORANGE)
    textbox(slide, 0.9, 0.8, 5, 0.35, "OPERATING SYSTEMS HACKATHON", 12, MINT, True)
    textbox(slide, 0.9, 1.65, 10.8, 1.7,
            "Mini Linux Shell\nwith Background Job Management", 34, WHITE, True,
            "Aptos Display")
    rect(slide, 0.9, 4.05, 1.2, 0.07, ORANGE)
    textbox(slide, 0.9, 4.45, 7.2, 0.55, "Team G-3  |  C11 + POSIX  |  Ubuntu Linux", 18, MINT)
    textbox(slide, 0.9, 6.75, 7.3, 0.3, "Processes made visible, manageable, and testable.", 13, WHITE)
    textbox(slide, 11.8, 6.9, 0.9, 0.25, "01 / 10", 10, MINT, True)

    split_slide(deck, 2, "The Problem", "Motivation", [
        ("One command, many OS concepts", "A shell translates text into process creation, descriptor wiring, and lifecycle decisions."),
        ("Failure paths matter", "Pipes, signals, and child exit states create resource and zombie risks."),
        ("Learning objective", "Build a small working system whose process behavior can be explained and verified."),
    ], "Project objective", [
        "Implement a compact Linux command interpreter.",
        "Keep ownership explicit: the shell tracks only children it creates.",
        "Demonstrate core POSIX interfaces without claiming full Bash compatibility.",
    ])

    split_slide(deck, 3, "Six Modules, One Shell", "Requirements", [
        ("01  Parse and execute", "Quoted words, syntax validation, fork/exec/wait."),
        ("02  Built-ins", "cd, pwd, help, exit, plus jobs and kill."),
        ("03  Background jobs", "&, process-group job records, Running/Done state."),
    ], "Process and I/O", [
        "04  Signals and cleanup: SIGCHLD reaping; bounded TERM-to-KILL.",
        "05  I/O: <, >, >>, and pipelines with pipe() and dup2().",
        "06  Errors: parser checks, syscall diagnostics, descriptor cleanup.",
    ])

    slide = new_slide(deck, 4, "Architecture", "Module boundaries")
    blocks = [
        (0.8, 2.15, 2.35, "Input loop", "main.c", TEAL),
        (3.55, 2.15, 2.35, "Parser", "parser.c", INK),
        (6.3, 2.15, 2.35, "Executor", "execute.c", ORANGE),
        (9.05, 2.15, 2.95, "Job manager", "jobs.c", TEAL),
    ]
    for x, y, w, heading, detail, color in blocks:
        rect(slide, x, y, w, 1.35, color, True)
        textbox(slide, x + 0.18, y + 0.2, w - 0.35, 0.4, heading, 17, WHITE, True)
        textbox(slide, x + 0.18, y + 0.72, w - 0.35, 0.3, detail, 12, WHITE)
    for x in (3.18, 5.93, 8.68):
        textbox(slide, x, 2.52, 0.35, 0.4, ">", 24, MUTED, True)
    rect(slide, 2.0, 4.45, 4.05, 1.2, PALE, True)
    textbox(slide, 2.3, 4.65, 3.5, 0.35, "builtins.c", 17, INK, True)
    textbox(slide, 2.3, 5.08, 3.5, 0.35, "Parent state + shell commands", 12, MUTED)
    rect(slide, 7.1, 4.45, 4.05, 1.2, MINT, True)
    textbox(slide, 7.4, 4.65, 3.5, 0.35, "shell.h", 17, TEAL, True)
    textbox(slide, 7.4, 5.08, 3.5, 0.35, "Shared Command, Pipeline, Job, Shell", 12, INK)
    textbox(slide, 0.9, 6.15, 10.9, 0.45,
            "The parser produces a pipeline; the executor creates processes; the job manager reaps owned children.",
            14, MUTED)

    split_slide(deck, 5, "Parsing and Built-ins", "Implementation", [
        ("Tokenize safely", "Quotes and backslash escapes group words; operators are recognized outside quotes."),
        ("Reject malformed input", "Incomplete pipes, missing paths, duplicate redirects, and misplaced & are errors."),
        ("Keep parent semantics", "Standalone cd and exit run in the shell process; pipeline built-ins run in children."),
    ], "Execution path", [
        "parse_line(line) -> Pipeline",
        "fork() one process per stage",
        "setpgid() for shared job ownership",
        "execvp() external commands",
        "waitpid() foreground processes",
    ])

    split_slide(deck, 6, "Pipelines and Redirection", "File descriptors", [
        ("Pipelines", "Each adjacent pair shares a pipe; children dup the correct ends onto stdin/stdout."),
        ("Input and output", "< opens read-only; > truncates or creates; >> appends or creates."),
        ("Descriptor ownership", "Parent and children close unused ends, including partial setup failures."),
    ], "Example", [
        "cat < input.txt | sort > sorted.txt",
        "\nPipe endpoints are connected first. Explicit redirection is then applied and overrides that stream for the stage.",
        "\nPipeline status is the final command's status.",
    ])

    split_slide(deck, 7, "Background Jobs", "Process groups", [
        ("Register once", "A pipeline started with & is recorded with job ID, PGID, member PIDs, and display text."),
        ("Reap in normal code", "SIGCHLD interrupts input; waitpid(WNOHANG) collects exits outside the signal handler."),
        ("Signal by job ID", "kill %N resolves only a job from this shell's table and targets its process group."),
    ], "Lifecycle", [
        "Running -> Done",
        "kill %N: SIGTERM",
        "Wait up to 1 second",
        "SIGKILL remaining group members",
        "waitpid() every recorded child",
    ])

    split_slide(deck, 8, "Safety and Error Handling", "Resource discipline", [
        ("Signal safety", "The handler does not print, allocate, or wait; child status is handled in the main flow."),
        ("No unrelated targets", "Only positive PGIDs stored for this shell's jobs are signaled; own shell group is guarded."),
        ("Failure cleanup", "Pipe descriptors start invalid and only opened endpoints are closed on setup failure."),
    ], "Graceful shutdown", [
        "On exit or EOF:",
        "1. Send SIGTERM to recorded active job groups.",
        "2. Poll/reap for a bounded grace period.",
        "3. Send SIGKILL to remaining owned groups.",
        "4. Wait for recorded children before freeing job state.",
    ])

    slide = new_slide(deck, 9, "Verified Results", "Testing")
    rect(slide, 0.85, 1.9, 3.1, 3.9, TEAL, True)
    textbox(slide, 1.15, 2.25, 2.5, 0.9, "10 / 10", 34, WHITE, True, "Aptos Display")
    textbox(slide, 1.15, 3.25, 2.35, 1.3, "integration tests\npassed", 21, WHITE, True)
    textbox(slide, 1.15, 5.15, 2.35, 0.35, "GCC 13.3.0  |  Ubuntu 24.04", 10, MINT)
    textbox(slide, 4.55, 2.05, 7.4, 0.42, "Coverage includes", 19, INK, True)
    textbox(slide, 4.55, 2.72, 7.3, 3.1,
            "Built-ins and parent cd\nQuoting, pipelines, and three redirection modes\nBackground state and safe job-scoped kill\nSIGTERM-resistant job and shutdown escalation\nDescriptor exhaustion and continued input\nSyntax errors and missing executable diagnostics",
            16, MUTED)
    textbox(slide, 4.55, 6.15, 7.5, 0.35, "Reproduce: make -B && make test", 13, TEAL, True)

    split_slide(deck, 10, "Impact and Next Steps", "Live demo", [
        ("Show", "Run a built-in, a pipeline, file redirection, then start, inspect, and terminate a job."),
        ("Explain", "Trace each command through parsing, fork/exec, descriptor wiring, and wait/reap."),
        ("Extend", "Add stopped-job tracking, fg/bg, expansion, or richer syntax as future lab modules."),
    ], "Scope, honestly stated", [
        "This is an educational shell, not Bash.",
        "No variable/command expansion, globbing, control flow, heredocs, fg/bg, or stopped-job table.",
        "Background marker is accepted only at the end of a pipeline.",
        "Commands execute with the user's normal permissions.",
    ])

    deck.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build_deck()